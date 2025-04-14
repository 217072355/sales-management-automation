import openpyxl
import shutil
import os

FILE_NAME = "sales.xlsx"

# Backup the original file before modifying
def backup_file():
    if os.path.exists(FILE_NAME):
        shutil.copy(FILE_NAME, "sales_backup.xlsx")
        print("Backup created as 'sales_backup.xlsx'")

# 1. Create a new Excel file if it doesn't exist
def create_excel_file():
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Sales Data"

    # Adding headers
    sheet.append(["Product", "Quantity", "Price per Unit"])

    # Saving file
    wb.save(FILE_NAME)
    print(f"Excel file '{FILE_NAME}' created successfully!")

# 2. Add a new product dynamically
def add_product():
    wb = openpyxl.load_workbook(FILE_NAME)
    sheet = wb.active

    product = input("Enter Product Name: ").strip()
    quantity = int(input("Enter Quantity: "))
    price = float(input("Enter Price per Unit: "))

    sheet.append([product, quantity, price])
    wb.save(FILE_NAME)
    print(f"Added {product} to {FILE_NAME}.")

# 3. Read data from the Excel file
def read_excel_file():
    wb = openpyxl.load_workbook(FILE_NAME)
    sheet = wb.active

    print("\nSales Data:")
    for row in sheet.iter_rows(values_only=True):
        print(row)

# 4. Search for a product
def search_product():
    wb = openpyxl.load_workbook(FILE_NAME)
    sheet = wb.active

    search_name = input("Enter product name to search: ").strip()
    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0].lower() == search_name.lower():
            print(f"Found: {row}")
            found = True
            break

    if not found:
        print("Product not found.")

# 5. Update the price of multiple products
def update_product_price():
    wb = openpyxl.load_workbook(FILE_NAME)
    sheet = wb.active

    while True:
        product_name = input("Enter product name to update (or 'done' to stop): ").strip()
        if product_name.lower() == "done":
            break

        new_price = float(input(f"Enter new price for {product_name}: "))

        updated = False
        for row in sheet.iter_rows(min_row=2, values_only=False):
            if row[0].value.lower() == product_name.lower():
                row[2].value = new_price
                updated = True
                print(f"Updated {product_name} price to {new_price}")

        if not updated:
            print(f"Product '{product_name}' not found.")

    wb.save(FILE_NAME)

# 6. Generate a sales report
def generate_sales_report():
    wb = openpyxl.load_workbook(FILE_NAME)
    sheet = wb.active

    total_revenue = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        total_revenue += row[1] * row[2]  # Quantity * Price

    report_sheet = wb.create_sheet("Sales Report")
    report_sheet.append(["Total Sales Revenue", total_revenue])
    wb.save(FILE_NAME)

    print(f"Sales report generated: Total Revenue = ${total_revenue}")

# Main Menu
def main():
    if not os.path.exists(FILE_NAME):
        create_excel_file()

    backup_file()

    while True:
        print("\nAdvanced Sales Management System")
        print("1. Add Product")
        print("2. View Sales Data")
        print("3. Search Product")
        print("4. Update Product Price")
        print("5. Generate Sales Report")
        print("6. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_product()
        elif choice == "2":
            read_excel_file()
        elif choice == "3":
            search_product()
        elif choice == "4":
            update_product_price()
        elif choice == "5":
            generate_sales_report()
        elif choice == "6":
            print("Exiting... Goodbye!")
            break
        else:
            print("Invalid choice. Please select again.")

# Run the program
main()
